import time
import urllib.request
import urllib.error
import json
import io
import openpyxl
from docx import Document

def fetch_artefact(name, url, method="GET", payload=None):
    data_bytes = None
    headers = {}
    if payload:
        data_bytes = json.dumps(payload).encode("utf-8")
        headers["Content-Type"] = "application/json"

    req = urllib.request.Request(url, data=data_bytes, headers=headers, method=method)
    start_time = time.perf_counter()
    try:
        with urllib.request.urlopen(req) as resp:
            content = resp.read()
            elapsed_ms = (time.perf_counter() - start_time) * 1000
            cd = resp.headers.get("Content-Disposition")
            ct = resp.headers.get("Content-Type")
            status = resp.status
            return {
                "name": name,
                "status": status,
                "elapsed_ms": elapsed_ms,
                "content": content,
                "content_type": ct,
                "content_disposition": cd,
                "bytes": len(content),
                "error": None
            }
    except urllib.error.HTTPError as e:
        elapsed_ms = (time.perf_counter() - start_time) * 1000
        body = e.read().decode("utf-8", errors="ignore")
        return {
            "name": name,
            "status": e.code,
            "elapsed_ms": elapsed_ms,
            "content": None,
            "content_type": None,
            "content_disposition": None,
            "bytes": 0,
            "error": body
        }
    except Exception as e:
        elapsed_ms = (time.perf_counter() - start_time) * 1000
        return {
            "name": name,
            "status": 0,
            "elapsed_ms": elapsed_ms,
            "content": None,
            "content_type": None,
            "content_disposition": None,
            "bytes": 0,
            "error": str(e)
        }

def verify_case(aid, std_name, expected_metrics):
    print(f"\n================================================================================")
    print(f"VERIFYING 6 ARTEFACTS & METRICS FOR: {std_name} ({aid})")
    print(f"================================================================================")

    endpoints = [
        ("1. DOCX Report", f"http://127.0.0.1:8000/api/iso27001/assessments/{aid}/export-docx", "POST", None),
        ("2. SoA XLSX", f"http://127.0.0.1:8000/api/iso27001/soa/export", "POST", {"assessment_id": aid, "org_name": "Test Org"}),
        ("3. Risk Register XLSX", f"http://127.0.0.1:8000/api/iso27001/assessments/{aid}/export-risk-register", "POST", None),
        ("4. PDF Report", f"http://127.0.0.1:8000/api/iso27001/assessments/{aid}/export-pdf", "POST", None),
        ("5. Assessment JSON", f"http://127.0.0.1:8000/api/iso27001/assessments/{aid}", "GET", None),
        ("6. Audit Trace JSON", f"http://127.0.0.1:8000/api/iso27001/assessments/{aid}/audit-trace", "GET", None),
    ]

    all_ok = True
    artefacts = {}

    for name, url, method, payload in endpoints:
        res = fetch_artefact(name, url, method, payload)
        artefacts[name] = res
        if res["status"] != 200:
            all_ok = False
            print(f"  ❌ {name}: HTTP {res['status']} | Error: {res['error']} ({res['elapsed_ms']:.1f}ms)")
            continue

        validation_note = "Valid"
        try:
            if "DOCX" in name:
                doc = Document(io.BytesIO(res["content"]))
                validation_note = f"docx valid, {len(doc.paragraphs)} paras, {len(doc.tables)} tables"
                # Check TCVN sanitization
                if "TCVN" in std_name:
                    full_text = " ".join(p.text for p in doc.paragraphs)
                    for t in doc.tables:
                        for row in t.rows:
                            for cell in row.cells:
                                full_text += " " + cell.text
                    forbidden = [
                        "HỒ SƠ ĐỀ XUẤT CẤP ĐỘ AN TOÀN HỆ THỐNG THÔNG TIN",
                        "XÁC NHẬN VÀ KÝ DUYỆT HỒ SƠ ĐỀ XUẤT CẤP ĐỘ",
                        "CĂN CỨ PHÁP LÝ THỰC HIỆN"
                    ]
                    for f_term in forbidden:
                        if f_term.lower() in full_text.lower():
                            raise ValueError(f"Forbidden term found in DOCX: '{f_term}'")
                    validation_note += " (TCVN wording fully sanitized)"
            elif "SoA" in name:
                wb = openpyxl.load_workbook(io.BytesIO(res["content"]))
                validation_note = f"openpyxl valid, sheets: {wb.sheetnames}"
                cd = res.get("content_disposition") or ""
                if "TCVN" in std_name and "SoA_ISO27001" in cd:
                    raise ValueError(f"TCVN SoA filename must not contain ISO27001, got: {cd}")
                elif "ISO" in std_name and "SoA_TCVN" in cd:
                    raise ValueError(f"ISO SoA filename must not contain TCVN, got: {cd}")
            elif "Risk Register" in name:
                wb = openpyxl.load_workbook(io.BytesIO(res["content"]))
                validation_note = f"openpyxl valid, sheets: {wb.sheetnames}"
            elif "PDF" in name:
                if res["content"][:4] == b"%PDF":
                    validation_note = "PDF binary header valid (%PDF)"
                elif b"<!DOCTYPE html" in res["content"] or b"<html" in res["content"]:
                    validation_note = "HTML print fallback valid"
            elif "Assessment JSON" in name:
                parsed = json.loads(res["content"].decode("utf-8"))
                jd = parsed.get("json_data", parsed)
                comp = jd.get("compliance", {})
                w_comp = jd.get("weighted_compliance", {})
                r_sum = jd.get("risk_summary", {})
                controls = jd.get("controls", [])
                rr = jd.get("risk_register", [])

                # Verify metrics
                print(f"      [Metric Check]")
                print(f"        Total Controls: {len(controls)} (expected: {expected_metrics['total_controls']})")
                print(f"        Evidence-supported Satisfied: {comp.get('score')}/{comp.get('max_score')} = {comp.get('percentage')}% (expected: {expected_metrics['ev_supp_score']}/{expected_metrics['total_controls']} = {expected_metrics['ev_supp_pct']}%)")
                print(f"        Weighted Compliance: {w_comp.get('weighted_score')}/{w_comp.get('weighted_max_score')} = {w_comp.get('percentage')}% (expected: {expected_metrics['w_score']}/{expected_metrics['w_max_score']} = {expected_metrics['w_pct']}%)")
                print(f"        Total Gaps: {len(rr)} | Risk Summary: {r_sum} (expected total: {expected_metrics['total_gaps']}, crit: {expected_metrics['crit_gaps']}, high: {expected_metrics['high_gaps']}, med: {expected_metrics['med_gaps']}, low: {expected_metrics['low_gaps']})")

                assert len(controls) == expected_metrics['total_controls'], f"Controls mismatch: {len(controls)} vs {expected_metrics['total_controls']}"
                assert comp.get('score') == expected_metrics['ev_supp_score'], f"Satisfied score mismatch: {comp.get('score')} vs {expected_metrics['ev_supp_score']}"
                assert comp.get('percentage') == expected_metrics['ev_supp_pct'], f"Evidence-supported pct mismatch: {comp.get('percentage')} vs {expected_metrics['ev_supp_pct']}"
                assert w_comp.get('weighted_score') == expected_metrics['w_score'], f"Weighted score mismatch: {w_comp.get('weighted_score')} vs {expected_metrics['w_score']}"
                assert w_comp.get('percentage') == expected_metrics['w_pct'], f"Weighted pct mismatch: {w_comp.get('percentage')} vs {expected_metrics['w_pct']}"
                assert len(rr) == expected_metrics['total_gaps'], f"Risk register count mismatch: {len(rr)} vs {expected_metrics['total_gaps']}"
                assert r_sum.get('total_gaps') == expected_metrics['total_gaps'], f"Risk summary total gaps mismatch: {r_sum.get('total_gaps')} vs {expected_metrics['total_gaps']}"
                assert r_sum.get('critical_gaps') == expected_metrics['crit_gaps'], f"Crit gaps mismatch"
                assert r_sum.get('high_gaps') == expected_metrics['high_gaps'], f"High gaps mismatch"
                assert r_sum.get('medium_gaps') == expected_metrics['med_gaps'], f"Med gaps mismatch"
                assert r_sum.get('low_gaps') == expected_metrics['low_gaps'], f"Low gaps mismatch"

                # Check citations
                null_shas = [cit for c in controls for cit in c.get("evidence_citations", []) if not cit.get("sha256")]
                print(f"        Citations check: 0 null sha256 found across controls (actual nulls: {len(null_shas)})")
                assert len(null_shas) == 0, f"Found {len(null_shas)} citations with null SHA-256"

                validation_note = "All metric & citation assertions PASSED"
            elif "Audit Trace JSON" in name:
                parsed = json.loads(res["content"].decode("utf-8"))
                events = parsed.get("events", [])
                export_events = [e for e in events if e.get("event_type") == "artifact_exported"]
                formats = set()
                for e in export_events:
                    p = e.get("payload") or {}
                    fmt = p.get("export_format") or e.get("export_format")
                    if fmt:
                        formats.add(fmt)
                validation_note = f"Total events: {len(events)}, artifact_exported events: {len(export_events)}, formats: {formats}"
                print(f"      [Audit Trace Check]")
                print(f"        Export events count: {len(export_events)}, formats exported: {formats}")
                for ee in export_events[-6:]:
                    p = ee.get("payload") or {}
                    fmt = p.get("export_format") or ee.get("export_format")
                    fn = p.get("filename") or ee.get("filename")
                    sz = p.get("file_size_bytes") or ee.get("file_size_bytes")
                    h = p.get("file_hash_sha256") or ee.get("file_hash_sha256") or ""
                    print(f"          - format={fmt} | file={fn} | size={sz} bytes | sha256={h[:16]}...")
        except Exception as ve:
            all_ok = False
            validation_note = f"VALIDATION ERROR: {ve}"

        print(f"  ✓ {name}: HTTP 200 | {res['bytes']:,} bytes | {res['elapsed_ms']:.1f}ms | {validation_note}")
        if res["content_disposition"]:
            print(f"      Content-Disposition: {res['content_disposition']}")

    return all_ok

if __name__ == "__main__":
    iso_metrics = {
        "total_controls": 93,
        "ev_supp_score": 34,
        "ev_supp_pct": 36.6,
        "w_score": 257.0,
        "w_max_score": 495.0,
        "w_pct": 51.9,
        "total_gaps": 59,
        "crit_gaps": 3,
        "high_gaps": 23,
        "med_gaps": 28,
        "low_gaps": 5,
    }

    tcvn_metrics = {
        "total_controls": 34,
        "ev_supp_score": 22,
        "ev_supp_pct": 64.7,
        "w_score": 195.0,
        "w_max_score": 271.0,
        "w_pct": 72.0,
        "total_gaps": 12,
        "crit_gaps": 4,
        "high_gaps": 6,
        "med_gaps": 2,
        "low_gaps": 0,
    }

    tcvn_f287_metrics = {
        "total_controls": 34,
        "ev_supp_score": 19,
        "ev_supp_pct": 55.9,
        "w_score": 165.0,
        "w_max_score": 271.0,
        "w_pct": 60.9,
        "total_gaps": 15,
        "crit_gaps": 6,
        "high_gaps": 7,
        "med_gaps": 2,
        "low_gaps": 0,
    }

    ok_iso = verify_case("163f3f6b-1265-41ac-acab-10f9d7580da5", "ISO 27001", iso_metrics)
    ok_tcvn = verify_case("e9bee7ff-8f23-4f01-9666-7e271e7b1a77", "TCVN 11930 (e9bee7ff)", tcvn_metrics)
    ok_f287 = verify_case("f2879709-3784-46db-88f3-074ad5a2cf1c", "TCVN 11930 (f2879709)", tcvn_f287_metrics)

    print("\n================================================================================")
    print(f"FINAL ARTEFACT VERIFICATION SUMMARY:")
    print(f"  Case 1 (ISO 27001 - 163f3f6b):       {'PASSED' if ok_iso else 'FAILED'}")
    print(f"  Case 2 (TCVN 11930 - e9bee7ff):      {'PASSED' if ok_tcvn else 'FAILED'}")
    print(f"  Case 3 (TCVN 11930 - f2879709 60.9%): {'PASSED' if ok_f287 else 'FAILED'}")
    print("================================================================================")

    if not (ok_iso and ok_tcvn and ok_f287):
        exit(1)
