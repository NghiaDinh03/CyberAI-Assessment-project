"""Risk Register .xlsx Exporter for CyberAI Assessment Platform.

Generates an IT Audit Risk Register spreadsheet with L x I quantitative risk scoring,
SOC severity badges, treatment plans, and ownership assignments.
"""

from __future__ import annotations

import io
import json
import logging
import os
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ── Color Palettes & Styles ──────────────────────────────────────────
_HEADER_FILL = PatternFill(start_color="0A2540", end_color="0A2540", fill_type="solid")
_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
_META_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
_META_FONT = Font(name="Calibri", size=11, bold=True, color="0A2540")
_BODY_FONT = Font(name="Calibri", size=10)
_BOLD_FONT = Font(name="Calibri", size=10, bold=True)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)

_SEV_FILLS = {
    "critical": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "high": PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid"),
    "medium": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"),
    "low": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
}

_SEV_FONTS = {
    "critical": Font(name="Calibri", size=10, bold=True, color="991B1B"),
    "high": Font(name="Calibri", size=10, bold=True, color="9A3412"),
    "medium": Font(name="Calibri", size=10, bold=True, color="854D0E"),
    "low": Font(name="Calibri", size=10, bold=True, color="334155"),
}


def generate_risk_register_xlsx(
    assessment_id: Optional[str] = None,
    assessment_data: Optional[Dict[str, Any]] = None,
    org_name: str = "",
) -> bytes:
    from schemas.assessment_schema import UnifiedAssessmentResult

    data = assessment_data or {}
    if assessment_id and not data:
        data_path = os.getenv("DATA_PATH", "./data")
        fpath = os.path.join(data_path, "assessments", f"{assessment_id}.json")
        if os.path.exists(fpath):
            try:
                with open(fpath, "r", encoding="utf-8") as f:
                    data = json.load(f)
            except Exception as e:
                logger.warning(f"Failed to load assessment {assessment_id}: {e}")

    if data:
        if isinstance(data, UnifiedAssessmentResult):
            validated = data
        else:
            validated = UnifiedAssessmentResult.model_validate(data)

        aid_str = validated.assessment_id
        run_id = validated.run_id
        code_version = validated.code_version
        raw_std = validated.standard
        std_display = raw_std.get("name") if isinstance(raw_std, dict) else str(raw_std)
        eval_date = validated.created_at[:10] if validated.created_at else datetime.now(timezone.utc).strftime("%Y-%m-%d")
        display_org = org_name or (validated.organization.get("name") if isinstance(validated.organization, dict) else "") or "Doanh nghiệp"
        weighted_comp_pct = None
        if hasattr(validated, "weighted_compliance") and validated.weighted_compliance:
            weighted_comp_pct = validated.weighted_compliance.percentage
        raw_risks = validated.risk_register or validated.top_gaps
        if not raw_risks and validated.controls:
            valid_risk_v = {"partial", "partially_satisfied", "missing", "not_evidenced", "needs_expert_review", "not_satisfied"}
            raw_risks = [
                c.model_dump() for c in validated.controls
                if c.assessment_verdict in valid_risk_v
            ]
    else:
        aid_str = assessment_id or "N/A"
        run_id = "run_default"
        code_version = "v1.2.0-verdict"
        std_display = "ISO 27001:2022"
        eval_date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        display_org = org_name or "Doanh nghiệp"
        raw_risks = []
        weighted_comp_pct = None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk Register"
    ws.views.sheetView[0].showGridLines = True

    # 1. Title & Meta Banner
    ws.merge_cells("A1:O1")
    t_cell = ws["A1"]
    t_cell.value = f"SỔ ĐĂNG KÝ RỦI RO AN TOÀN THÔNG TIN (RISK REGISTER) — {display_org.upper()}"
    t_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t_cell.fill = PatternFill(start_color="0A2540", end_color="0A2540", fill_type="solid")
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Row 2: Metadata Info
    ws.merge_cells("A2:O2")
    m_cell = ws["A2"]
    comp_str = f"   |   Tuân thủ: {weighted_comp_pct:.1f}%" if weighted_comp_pct is not None else ""
    m_cell.value = (
        f"Assessment ID: {aid_str}   |   Run ID: {run_id}   |   Code Version: {code_version}   |   "
        f"Tiêu chuẩn: {std_display}{comp_str}   |   Ngày lập: {eval_date}   |   Nền tảng: CyberAI Platform"
    )
    m_cell.font = Font(name="Calibri", size=10, italic=True, color="334155")
    m_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    m_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 24

    # Row 3: Headers
    headers = [
        ("STT", 6),
        ("Mã Control", 13),
        ("Tên Control / Tiêu Chuẩn", 26),
        ("Nhóm Kiểm Soát", 20),
        ("Khoảng Trống An Ninh (GAP)", 36),
        ("Mức Độ", 13),
        ("Khả Năng (L: 1-4)", 15),
        ("Tác Động (I: 1-4)", 15),
        ("Điểm Rủi Ro (L×I)", 16),
        ("Cơ sở Đánh giá", 18),
        ("Khuyến Nghị Khắc Phục", 36),
        ("Phương Án Xử Lý", 18),
        ("Thời Hạn", 14),
        ("Chủ Sở Hữu", 18),
        ("Trạng Thái", 14),
    ]

    for col_idx, (h_title, col_width) in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h_title)
        cell.font = _HEADER_FONT
        cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    ws.row_dimensions[3].height = 28

    # Extract risk rows with strict filtering
    raw_risks = raw_risks or []
    valid_risk_verdicts = {"partial", "partially_satisfied", "missing", "not_evidenced", "needs_expert_review", "not_satisfied"}

    ctrl_map = {}
    if data and 'validated' in locals() and hasattr(validated, "controls") and validated.controls:
        ctrl_map = {c.control_id: c for c in validated.controls}

    filtered_risks = []
    for r_item in raw_risks:
        cid = str(r_item.get("control_id") or r_item.get("id") or "")
        ctrl_obj = ctrl_map.get(cid)
        v = str(r_item.get("assessment_verdict") or r_item.get("verdict") or (ctrl_obj.assessment_verdict if ctrl_obj else "")).lower()
        u = str(r_item.get("user_declaration") or (ctrl_obj.user_declaration if ctrl_obj else "")).lower()
        if v == "satisfied":
            continue
        if v and v not in valid_risk_verdicts:
            continue
        filtered_risks.append(r_item)

    from services.risk_scoring import calculate_risk_score

    # Sort by risk_score desc
    def _score(r):
        return int(r.get("risk_score") or (int(r.get("likelihood", 3)) * int(r.get("impact", 3))))

    sorted_risks = sorted(filtered_risks, key=_score, reverse=True)

    current_row = 4
    if not sorted_risks:
        ws.merge_cells("A4:O4")
        empty_cell = ws["A4"]
        empty_cell.value = "Không có mục rủi ro được tạo từ kết quả hiện tại (Không có rủi ro cần đưa vào báo cáo)"
        empty_cell.font = Font(name="Calibri", size=11, italic=True, color="64748B")
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        empty_cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        for c in range(1, 16):
            ws.cell(row=4, column=c).border = _THIN_BORDER
        ws.row_dimensions[4].height = 28
    else:
        for idx, r_item in enumerate(sorted_risks, 1):
            ctrl_id = str(r_item.get("control_id") or r_item.get("id") or "N/A")
            ctrl_label = str(r_item.get("label") or ctrl_id)
            cat = str(r_item.get("category") or "General")
            gap = str(r_item.get("gap") or "Khoảng trống kiểm soát cần bổ sung bằng chứng hoặc chính sách.")
            sev = str(r_item.get("severity") or r_item.get("risk_severity") or "medium").lower()
            l_val = int(r_item.get("likelihood") or (4 if sev == "critical" else 3 if sev == "high" else 2 if sev == "medium" else 1))
            i_val = int(r_item.get("impact") or (4 if sev == "critical" else 3 if sev == "high" else 2 if sev == "medium" else 1))
            # Ensure values conform to official 1-4 range
            l_val = min(4, max(1, l_val))
            i_val = min(4, max(1, i_val))
            r_score = int(r_item.get("risk_score") or calculate_risk_score(l_val, i_val))
            if r_score > 16:
                r_score = calculate_risk_score(l_val, i_val)
            basis = str(r_item.get("risk_assessment_basis") or "rule_based")
            rec = str(r_item.get("recommendation") or "Ban hành quy trình và thực thi biện pháp kỹ thuật.")
            treatment = "Giảm thiểu (Mitigate)" if r_score >= 12 else "Chấp nhận có điều kiện" if r_score <= 4 else "Kiểm soát bổ sung"
            timeline = str(r_item.get("timeline") or ("0-30 ngày" if sev in ("critical", "high") else "1-3 tháng"))
            owner = "CISO / IT Security" if sev in ("critical", "high") else "Hạ tầng / SysAdmin"
            status = "Đang mở (Open)"

            row_vals = [
                (idx, _CENTER, _BODY_FONT, None),
                (ctrl_id, _CENTER, _BOLD_FONT, None),
                (ctrl_label, _LEFT, _BODY_FONT, None),
                (cat, _LEFT, _BODY_FONT, None),
                (gap, _LEFT, _BODY_FONT, None),
                (sev.upper(), _CENTER, _SEV_FONTS.get(sev, _BOLD_FONT), _SEV_FILLS.get(sev)),
                (l_val, _CENTER, _BODY_FONT, None),
                (i_val, _CENTER, _BODY_FONT, None),
                (r_score, _CENTER, _SEV_FONTS.get(sev, _BOLD_FONT), _SEV_FILLS.get(sev)),
                (basis, _CENTER, _BODY_FONT, None),
                (rec, _LEFT, _BODY_FONT, None),
                (treatment, _CENTER, _BODY_FONT, None),
                (timeline, _CENTER, _BODY_FONT, None),
                (owner, _CENTER, _BODY_FONT, None),
                (status, _CENTER, _BODY_FONT, None),
            ]

            for col_idx, (val, align, font_style, fill_style) in enumerate(row_vals, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.alignment = align
                cell.font = font_style
                cell.border = _THIN_BORDER
                if fill_style:
                    cell.fill = fill_style

            ws.row_dimensions[current_row].height = 24
            current_row += 1

    # Freeze panes below header
    ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
