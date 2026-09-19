"""Unit and Integration Tests for Multi-Format Evidence Parser and Batch Ingest API."""

import io
import json
import pytest
from services.evidence_parser import (
    parse_evidence_file,
    parse_text_file,
    parse_docx_file,
    parse_pdf_file,
    parse_image_file,
    MAX_EVIDENCE_SIZE_BYTES,
)


def test_parse_text_encodings():
    # 1. UTF-8
    utf8_bytes = "Chính sách an toàn thông tin TCVN 11930:2017".encode("utf-8")
    res = parse_evidence_file(utf8_bytes, "policy.txt")
    assert res["status"] == "success"
    assert "TCVN 11930:2017" in res["parsed_text"]
    assert res["char_count"] > 0

    # 2. UTF-8 with BOM
    bom_bytes = b"\xef\xbb\xbf" + "Log he thong Active Directory 10.140.0.3".encode("utf-8")
    res_bom = parse_evidence_file(bom_bytes, "ad_auth.log")
    assert res_bom["status"] == "success"
    assert "10.140.0.3" in res_bom["parsed_text"]

    # 3. CSV log
    csv_bytes = "timestamp,user,action,status\n2026-08-26,admin,login,success".encode("utf-8")
    res_csv = parse_evidence_file(csv_bytes, "audit.csv")
    assert res_csv["status"] == "success"
    assert "admin" in res_csv["parsed_text"]

    # 4. JSON config
    json_bytes = json.dumps({"firewall": "NGFW", "version": "10.2", "dmz": True}).encode("utf-8")
    res_json = parse_evidence_file(json_bytes, "firewall_cfg.json")
    assert res_json["status"] == "success"
    assert "NGFW" in res_json["parsed_text"]


def test_parse_unsupported_format():
    exe_bytes = b"MZ\x90\x00\x03\x00\x00\x00"
    res = parse_evidence_file(exe_bytes, "malicious_payload.exe")
    assert res["status"] == "unsupported"
    assert res["error_code"] == "UNSUPPORTED_FORMAT"
    assert res["char_count"] == 0


def test_parse_oversized_file():
    oversized_bytes = b"A" * (MAX_EVIDENCE_SIZE_BYTES + 1024)
    res = parse_evidence_file(oversized_bytes, "huge_dump.log")
    assert res["status"] == "failed"
    assert res["error_code"] == "FILE_TOO_LARGE"


def test_parse_docx_mock():
    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph("Quy chế quản lý sao lưu dữ liệu máy chủ")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Máy chủ"
        table.cell(0, 1).text = "Chu kỳ Backup"
        table.cell(1, 0).text = "EOFFICE-DB"
        table.cell(1, 1).text = "Daily Incremental"
        stream = io.BytesIO()
        doc.save(stream)
        docx_bytes = stream.getvalue()

        res = parse_evidence_file(docx_bytes, "backup_policy.docx")
        assert res["status"] == "success"
        assert "sao lưu" in res["parsed_text"]
        assert "EOFFICE-DB" in res["parsed_text"]
    except ImportError:
        pytest.skip("python-docx not installed in test environment")


def test_parse_image_metadata_and_ocr_fallback():
    try:
        from PIL import Image
        img = Image.new("RGB", (200, 100), color=(255, 255, 255))
        stream = io.BytesIO()
        img.save(stream, format="PNG")
        png_bytes = stream.getvalue()

        res = parse_evidence_file(png_bytes, "network_topology.png")
        assert res["status"] == "success"
        assert "200x100" in res["parsed_text"]
    except ImportError:
        pytest.skip("PIL not installed in test environment")


def test_cross_verify_controls():
    from services.evidence_fact_extractor import EvidenceFactExtractor

    # Evidence 1: Host log with Windows Server 2008 R2 (EOL) and Hotfix: N/A
    log_text = """
    Host Name:                 SRV-APP-01
    OS Name:                   Microsoft Windows Server 2008 R2 Standard
    OS Version:                6.1.7601 Service Pack 1 Build 7601
    Hotfix(s):                 N/A
    Domain:                    thuducpp.evn.vn
    """
    card1 = EvidenceFactExtractor.extract_facts(log_text, "10.140.0.11.txt", use_llm=False)

    # Evidence 2: Policy doc with approval
    policy_text = """
    CÔNG TY TNHH MTV NHIỆT ĐIỆN THỦ ĐỨC
    QUY CHẾ AN TOÀN THÔNG TIN NỘI BỘ
    Ngày ban hành: 15/01/2026. Người phê duyệt: Giám đốc Công ty.
    """
    card2 = EvidenceFactExtractor.extract_facts(policy_text, "chinh_sach_attt.docx", use_llm=False)

    # Self-attestation list: user checks SV.07, A.5.1, and A.5.7
    self_attested = ["SV.07", "A.5.1", "A.5.7"]
    res = EvidenceFactExtractor.cross_verify_controls(self_attested, [card1, card2])

    assert res["total_attested"] == 3
    # SV.07 has contradiction because OS is EOL and Hotfix is N/A
    assert any(c["control_id"] == "SV.07" for c in res["contradiction_gaps"])
    # A.5.1 is verified because policy document has title and approval
    assert any(c["control_id"] == "A.5.1" for c in res["verified_satisfied"])
    # A.5.7 (Threat intelligence) has no evidence in files -> unverified oversight
    assert any(c["control_id"] == "A.5.7" for c in res["unverified_oversights"])

    # Verify formatting output
    formatted = EvidenceFactExtractor.format_for_auditor([card1, card2], self_attested_controls=self_attested)
    assert "MÂU THUẪN NGHIÊM TRỌNG" in formatted
    assert "CHƯA ĐƯỢC XÁC THỰC" in formatted
    assert "Tự khai báo - Không có log đối chứng" in formatted
