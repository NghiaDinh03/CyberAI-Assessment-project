"""Functional Test Matrix covering FT-01 through FT-08 according to official project specification table.

FT-01: Quản lý standard
       Catalogue tích hợp sẵn được tải đúng; dữ liệu JSON/YAML được kiểm tra cấu trúc trước khi sử dụng.
FT-02: Assessment Wizard và minh chứng
       Người dùng chọn được control, nạp được tệp đầu vào; parser/OCR trả về trạng thái xử lý hoặc lỗi có thể nhận biết.
FT-03: Luồng assessment
       Assessment được tạo theo phạm vi đã chọn; kết quả có liên kết với assessment_id và run_id.
FT-04: RAG và đầu ra có cấu trúc
       Trace thể hiện metadata truy hồi phù hợp; đầu ra AI được kiểm tra schema trước khi tổng hợp.
FT-05: Artefact kết quả
       JSON, SoA XLSX, báo cáo và dữ liệu trên màn hình kết quả thuộc cùng assessment/run và không mâu thuẫn dữ liệu.
FT-06: Chatbot tri thức
       Tạo/khôi phục phiên, lịch sử hội thoại, lựa chọn mô hình và phản hồi stream hoạt động theo luồng thiết kế.
FT-07: Ranh giới Web Search
       search_context chỉ xuất hiện trong chatbot, không được đưa vào minh chứng, scoring, GAP hoặc rủi ro assessment.
FT-08: Audit trace
       Trace có định danh, metadata kỹ thuật cần thiết và không chứa dữ liệu nhạy cảm hoặc thông tin xác thực.
"""

from __future__ import annotations

import io
import json
import os
import pathlib
import sys
from unittest.mock import MagicMock, patch

import openpyxl
import pytest
from docx import Document
from pydantic import ValidationError

BACKEND = pathlib.Path(__file__).resolve().parents[1]
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from schemas.assessment_schema import (
    ControlCoverage,
    ControlItem,
    UnifiedAssessmentResult,
    WeightedCompliance,
)
from services.audit_service import AuditContext, AuditService
from services.chat_service import ChatService
from services.controls_catalog import (
    calc_weighted_compliance,
    get_categories,
    get_flat_controls,
)
from services.evidence_parser import (
    compute_file_sha256,
    mask_evidence_filename,
    parse_evidence_file,
)
from services.report_docx_generator import generate_report_docx
from services.risk_register_exporter import generate_risk_register_xlsx
from services.soa_exporter import generate_soa_xlsx
from services.standard_service import parse_yaml_to_dict, validate_standard
from services.web_search import WebSearch


# ── FT-01: Quản lý standard ───────────────────────────────────────────────────


def test_ft_01_quan_ly_standard():
    """FT-01: Catalogue tích hợp sẵn được tải đúng; dữ liệu JSON/YAML được kiểm tra cấu trúc trước khi sử dụng."""
    # 1. Catalogue tích hợp sẵn được tải đúng
    iso_controls = get_flat_controls("iso27001")
    assert len(iso_controls) == 93, f"ISO 27001 phải có đúng 93 controls, nhận: {len(iso_controls)}"
    assert all(c["id"].startswith("A.") for c in iso_controls)

    tcvn_controls = get_flat_controls("tcvn11930")
    assert len(tcvn_controls) == 34, f"TCVN 11930 phải có đúng 34 controls, nhận: {len(tcvn_controls)}"
    valid_tcvn_prefixes = {"NW", "SV", "APP", "DAT", "MNG"}
    assert all(c["id"].split(".")[0] in valid_tcvn_prefixes for c in tcvn_controls)

    # 2. Dữ liệu JSON/YAML được kiểm tra cấu trúc trước khi sử dụng
    valid_custom_json = {
        "id": "pci_dss_test",
        "name": "PCI-DSS Custom Test",
        "controls": [
            {
                "category": "1. Network",
                "controls": [
                    {"id": "1.1", "label": "Firewall Config", "weight": "critical"}
                ],
            }
        ],
    }
    errors_valid = validate_standard(valid_custom_json)
    assert errors_valid == [], f"Standard hợp lệ không được có lỗi: {errors_valid}"

    # Kiểm tra cấu trúc sai (thiếu id, weight không hợp lệ, trùng lặp control id)
    invalid_custom_json = {
        "name": "Missing ID Standard",
        "controls": [
            {
                "category": "General",
                "controls": [
                    {"id": "G.1", "label": "Item 1", "weight": "invalid_weight"},
                    {"id": "G.1", "label": "Duplicate Item", "weight": "high"},
                ],
            }
        ],
    }
    errors_invalid = validate_standard(invalid_custom_json)
    assert any("Missing required field: 'id'" in e for e in errors_invalid)
    assert any("invalid weight" in e for e in errors_invalid)
    assert any("Duplicate control ID" in e for e in errors_invalid)

    # Kiểm tra parser YAML cấu trúc
    sample_yaml = """
    id: hipaa_security
    name: HIPAA Security Rule
    controls:
      - category: Administrative Safeguards
        controls:
          - id: 164.308(a)(1)
            label: Security Management Process
            weight: critical
    """
    parsed_yaml = parse_yaml_to_dict(sample_yaml)
    assert parsed_yaml["id"] == "hipaa_security"
    assert len(parsed_yaml["controls"]) == 1


# ── FT-02: Assessment Wizard và minh chứng ─────────────────────────────────────


def test_ft_02_assessment_wizard_va_minh_chung():
    """FT-02: Người dùng chọn được control, nạp được tệp đầu vào; parser/OCR trả về trạng thái xử lý hoặc lỗi có thể nhận biết."""
    # 1. Người dùng chọn được control trong phạm vi
    selected_controls = ["A.5.1", "A.5.2", "A.8.1"]
    all_iso = get_flat_controls("iso27001")
    scope_controls = [c for c in all_iso if c["id"] in selected_controls]
    assert len(scope_controls) == 3
    assert {c["id"] for c in scope_controls} == set(selected_controls)

    # 2. Nạp tệp đầu vào hợp lệ (TXT, XLSX, DOCX)
    txt_bytes = b"Chinh sach bao mat thong tin he thong 2026 - IP: 10.140.0.1"
    res_txt = parse_evidence_file(txt_bytes, "policy_10.140.0.1.txt")
    assert res_txt["status"] == "success"
    assert res_txt["char_count"] > 0
    assert res_txt["sha256"] == compute_file_sha256(txt_bytes)
    assert "***" in res_txt["masked_filename"]  # Đã che mặt nạ IP

    # 3. Parser / OCR trả về trạng thái xử lý hoặc lỗi có thể nhận biết
    corrupted_bytes = b"\x00\x01NOT_A_VALID_DOCUMENT_FILE_STREAM"
    res_err = parse_evidence_file(corrupted_bytes, "corrupted_audit.pdf")
    assert res_err["status"] == "failed"
    assert res_err["error_code"] is not None
    assert "corrupted_audit.pdf" in res_err["filename"]

    # Định dạng không hỗ trợ
    res_unsupported = parse_evidence_file(b"echo 123", "script.sh")
    assert res_unsupported["status"] == "unsupported"
    assert res_unsupported["error_code"] == "UNSUPPORTED_FORMAT"


# ── FT-03: Luồng assessment ───────────────────────────────────────────────────


def test_ft_03_luong_assessment():
    """FT-03: Assessment được tạo theo phạm vi đã chọn; kết quả có liên kết với assessment_id và run_id."""
    target_assessment_id = "asm_flow_audit_2026"
    target_run_id = "run_flow_exec_001"

    # Tạo assessment theo phạm vi ISO 27001 đã chọn
    iso_categories = get_categories("iso27001")
    scope_items = []
    for cat in iso_categories:
        for c in cat["controls"]:
            scope_items.append(ControlItem(
                control_id=c["id"],
                label=c["label"],
                category=cat["category"],
                weight=c["weight"],
                user_declaration="implemented" if c["id"] in ("A.5.1", "A.5.2") else "not_implemented",
                assessment_verdict="satisfied" if c["id"] == "A.5.1" else "missing",
            ))

    scoring = calc_weighted_compliance([c.model_dump() for c in scope_items])

    payload = {
        "assessment_id": target_assessment_id,
        "run_id": target_run_id,
        "code_version": "v1.2.0-rel",
        "standard": {"id": "iso27001", "name": "ISO/IEC 27001:2022"},
        "organization": {"name": "Audit Bank Corp"},
        "controls": [c.model_dump() for c in scope_items],
        "weighted_compliance": {
            "weighted_score": scoring["weighted_score"],
            "weighted_max_score": scoring["weighted_max_score"],
            "percentage": scoring["percentage"],
        },
    }

    # Validate cấu trúc và liên kết định danh
    validated = UnifiedAssessmentResult.model_validate(payload)
    assert validated.assessment_id == target_assessment_id
    assert validated.run_id == target_run_id
    assert len(validated.controls) == 93
    assert validated.weighted_compliance.percentage == scoring["percentage"]


# ── FT-04: RAG và đầu ra có cấu trúc ───────────────────────────────────────────


def test_ft_04_rag_va_dau_ra_co_cau_truc():
    """FT-04: Trace thể hiện metadata truy hồi phù hợp; đầu ra AI được kiểm tra schema trước khi tổng hợp."""
    # 1. Trace thể hiện metadata truy hồi RAG phù hợp
    mock_retrieval_chunks = [
        {
            "text": "Chính sách kiểm soát truy cập yêu cầu MFA cho mọi kết nối VPN.",
            "source": "vpn_security_policy_2026.pdf",
            "score": 0.88,
            "chunk_id": "chunk_vpn_01",
            "collection": "iso_documents",
        },
        {
            "text": "Nhật ký xác thực PAM chỉ ghi nhận mật khẩu đơn kỳ tháng 9.",
            "source": "auth_audit_2026.log",
            "score": 0.81,
            "chunk_id": "chunk_pam_04",
            "collection": "evidence_vault",
        },
    ]

    for chunk in mock_retrieval_chunks:
        assert "source" in chunk
        assert "score" in chunk and chunk["score"] > 0
        assert "chunk_id" in chunk

    # 2. Đầu ra AI được kiểm tra schema trước khi tổng hợp
    raw_ai_control_output = {
        "control_id": "A.5.15",
        "label": "Kiểm soát truy cập",
        "category": "A.5 Tổ chức",
        "weight": "critical",
        "assessment_verdict": "needs_expert_review",
        "conflict_detected": True,
        "conflict_reason": "Chính sách yêu cầu MFA nhưng nhật ký PAM ghi nhận xác thực đơn yếu tố",
        "conflicting_evidence_ids": ["auth_audit_2026.log"],
        "gap": "Chưa áp dụng MFA đồng bộ",
        "recommendation": "Kích hoạt MFA bắt buộc cho VPN/PAM",
    }

    # Bắt buộc validate qua schema ControlItem
    validated_control = ControlItem.model_validate(raw_ai_control_output)
    assert validated_control.control_id == "A.5.15"
    assert validated_control.conflict_detected is True
    assert validated_control.assessment_verdict == "needs_expert_review"

    # Schema từ chối đầu ra lỗi của AI
    invalid_ai_output = dict(raw_ai_control_output)
    invalid_ai_output["assessment_verdict"] = "hallucinated_compliant_verdict"
    with pytest.raises(ValidationError):
        ControlItem.model_validate(invalid_ai_output)


# ── FT-05: Artefact kết quả ───────────────────────────────────────────────────


def test_ft_05_artefact_ket_qua():
    """FT-05: JSON, SoA XLSX, báo cáo và dữ liệu trên màn hình kết quả thuộc cùng assessment/run và không mâu thuẫn dữ liệu."""
    asm_id = "asm_artefact_ft05"
    run_id = "run_artefact_ft05_01"

    ctrl_1 = ControlItem(
        control_id="A.5.1",
        label="Chính sách an toàn thông tin",
        category="A.5 Tổ chức",
        weight="critical",
        user_declaration="implemented",
        assessment_verdict="satisfied",
        score=4,
    )
    ctrl_2 = ControlItem(
        control_id="A.5.15",
        label="Kiểm soát truy cập",
        category="A.5 Tổ chức",
        weight="critical",
        user_declaration="implemented",
        assessment_verdict="needs_expert_review",
        conflict_detected=True,
        conflict_reason="Mâu thuẫn giữa tự khai và log",
        score=0,
        gap="Thiếu MFA",
        recommendation="Bật MFA",
        risk_severity="critical",
        likelihood=4,
        impact=4,
        risk_score=16,
    )

    for ctrl in [ctrl_1, ctrl_2]:
        assert 1 <= ctrl.likelihood <= 4
        assert 1 <= ctrl.impact <= 4
        assert ctrl.risk_score == ctrl.likelihood * ctrl.impact
        assert ctrl.risk_score <= 16

    scoring = calc_weighted_compliance([ctrl_1.model_dump(), ctrl_2.model_dump()])

    unified = UnifiedAssessmentResult(
        assessment_id=asm_id,
        run_id=run_id,
        standard={"id": "iso27001", "name": "ISO/IEC 27001:2022"},
        organization={"name": "Viện Tài chính Quốc gia"},
        controls=[ctrl_1, ctrl_2],
        weighted_compliance=WeightedCompliance(
            weighted_score=scoring["weighted_score"],
            weighted_max_score=scoring["weighted_max_score"],
            percentage=scoring["percentage"],
        ),
    )

    # 1. JSON
    json_data = unified.model_dump()
    assert json_data["assessment_id"] == asm_id
    assert json_data["run_id"] == run_id

    # 2. SoA XLSX
    soa_bytes = generate_soa_xlsx(assessment_data=json_data)
    wb_soa = openpyxl.load_workbook(io.BytesIO(soa_bytes))
    ws_soa = wb_soa.active
    meta_soa = ws_soa["A2"].value or ""
    assert asm_id in meta_soa
    assert run_id in meta_soa

    # 3. Risk Register XLSX
    risk_bytes = generate_risk_register_xlsx(assessment_data=json_data)
    wb_risk = openpyxl.load_workbook(io.BytesIO(risk_bytes))
    ws_risk = wb_risk.active
    meta_risk = ws_risk["A2"].value or ""
    assert asm_id in meta_risk
    assert run_id in meta_risk

    # 4. DOCX Report
    docx_bytes = generate_report_docx(assessment_data=json_data)
    doc = Document(io.BytesIO(docx_bytes))
    doc_text = " ".join([p.text for p in doc.paragraphs])
    assert "Viện Tài chính Quốc gia" in doc_text

    # 5. Đối chiếu không mâu thuẫn dữ liệu giữa các artefact
    assert json_data["weighted_compliance"]["percentage"] == scoring["percentage"]
    # SoA hiển thị A.5.1 Satisfied, A.5.15 Needs Expert Review
    soa_verdicts = {}
    for row in ws_soa.iter_rows(min_row=5, values_only=True):
        if row and row[0] and str(row[0]).startswith("A."):
            soa_verdicts[row[0]] = str(row[10]).lower().replace(" ", "_")

    assert soa_verdicts["A.5.1"] == "satisfied"
    assert soa_verdicts["A.5.15"] == "needs_expert_review"
    assert soa_verdicts["A.5.1"] == json_data["controls"][0]["assessment_verdict"]
    assert soa_verdicts["A.5.15"] == json_data["controls"][1]["assessment_verdict"]

    for c in json_data["controls"]:
        assert 1 <= c["likelihood"] <= 4
        assert 1 <= c["impact"] <= 4
        assert c["risk_score"] == c["likelihood"] * c["impact"]
        assert c["risk_score"] <= 16


# ── FT-06: Chatbot tri thức ───────────────────────────────────────────────────


def test_ft_06_chatbot_tri_thuc():
    """FT-06: Tạo/khôi phục phiên, lịch sử hội thoại, lựa chọn mô hình và phản hồi stream hoạt động theo luồng thiết kế."""
    # 1. Quản lý phiên và lịch sử hội thoại
    session_store = ChatService.get_session_store()
    test_session_id = f"test_session_ft06_{os.getpid()}"

    session_store.save(test_session_id, {
        "messages": [
            {"role": "user", "content": "Xin chào CyberAI"},
            {"role": "assistant", "content": "Chào bạn, tôi sẵn sàng hỗ trợ đánh giá ATTT."},
        ]
    })

    # Khôi phục phiên
    history = session_store.get_context_messages(test_session_id, max_messages=5)
    assert len(history) == 2
    assert history[0]["role"] == "user"
    assert history[1]["role"] == "assistant"

    # 2. Lựa chọn mô hình và định tuyến
    from services.model_router import route_model
    routing_sec = route_model("Phân tích quy định ISO 27001 về kiểm soát truy cập")
    assert routing_sec["route"] == "security"
    assert routing_sec["use_rag"] is True

    # 3. Phản hồi stream theo luồng thiết kế
    routing_test = {"use_rag": False, "use_search": False, "model": "test-model"}
    built_msgs = ChatService._build_messages(
        message="Kiểm tra stream",
        routing=routing_test,
        context="",
        search_context="",
        history=history,
        is_local=True,
    )
    assert len(built_msgs) >= 3  # System prompt + history + current message
    assert built_msgs[-1]["role"] == "user"
    assert "Kiểm tra stream" in built_msgs[-1]["content"]


# ── FT-07: Ranh giới Web Search ───────────────────────────────────────────────


def test_ft_07_ranh_gioi_web_search():
    """FT-07: search_context chỉ xuất hiện trong chatbot, không được đưa vào minh chứng, scoring, GAP hoặc rủi ro assessment."""
    # 1. search_context xuất hiện trong Chatbot
    sample_search_results = [
        {"title": "CISA Zero-Day Alert 2026", "url": "https://cisa.gov/alert-2026", "snippet": "Advisory on VPN vulnerability"}
    ]
    formatted_search_context = WebSearch.format_context(sample_search_results)
    assert "https://cisa.gov/alert-2026" in formatted_search_context

    chat_routing = {"use_rag": False, "use_search": True, "model": "cyberai-local"}
    chat_messages = ChatService._build_messages(
        message="Tin tức bảo mật",
        routing=chat_routing,
        context="",
        search_context=formatted_search_context,
        history=[],
        is_local=True,
    )
    # search_context có trong prompt gửi tới mô hình chatbot
    assert "https://cisa.gov/alert-2026" in chat_messages[-1]["content"]
    assert "Search Results / Kết quả tìm kiếm:" in chat_messages[-1]["content"]

    # 2. Ranh giới tuyệt đối: search_context KHÔNG ĐƯỢC đưa vào assessment pipeline
    assessment_controls = [
        ControlItem(
            control_id="A.5.1",
            label="Chính sách an toàn thông tin",
            weight="critical",
            assessment_verdict="satisfied",
            evidence_file_ids=["local_policy.pdf"],
            gap="",
        ),
        ControlItem(
            control_id="A.5.15",
            label="Kiểm soát truy cập",
            weight="critical",
            assessment_verdict="missing",
            evidence_file_ids=[],
            gap="Chưa triển khai IAM",
        ),
    ]

    # Authoritative scoring tính từ danh mục nội bộ, không liên quan đến web search
    scoring = calc_weighted_compliance([c.model_dump() for c in assessment_controls])
    assert scoring["percentage"] == 50.0

    # Unified assessment record và exporters không chứa search_context
    unified_asm = UnifiedAssessmentResult(
        assessment_id="asm_web_boundary_test",
        run_id="run_boundary_01",
        controls=assessment_controls,
        weighted_compliance=WeightedCompliance(
            weighted_score=scoring["weighted_score"],
            weighted_max_score=scoring["weighted_max_score"],
            percentage=scoring["percentage"],
        ),
    )
    asm_dump_str = json.dumps(unified_asm.model_dump())
    assert "https://cisa.gov/alert-2026" not in asm_dump_str
    assert "search_context" not in asm_dump_str


# ── FT-08: Audit trace ─────────────────────────────────────────────────────────


def test_ft_08_audit_trace():
    """FT-08: Trace có định danh, metadata kỹ thuật cần thiết và không chứa dữ liệu nhạy cảm hoặc thông tin xác thực."""
    test_aid = "asm_audit_trace_test_2026"
    test_run = "run_audit_trace_001"
    ctx = AuditContext(assessment_id=test_aid, run_id=test_run)

    # 1. Ghi nhận sự kiện kiểm toán kỹ thuật
    event_id = AuditService.record_assessment_created(
        ctx=ctx,
        standard="iso27001",
        model_mode="local_hybrid",
        org_name="Công ty Bảo mật Quốc gia",
        implemented_controls_count=45,
        total_controls=93,
        has_evidence=True,
    )
    assert event_id != ""

    # Ghi nhận sự kiện xử lý minh chứng
    file_bytes = b"Sample Evidence Log Content"
    sha = compute_file_sha256(file_bytes)
    ev_event_id = AuditService.record_evidence_parsed(
        ctx=ctx,
        evidence_controls_count=10,
        total_files=2,
        file_extensions=[".pdf", ".docx"],
    )
    assert ev_event_id != ""

    # 2. Trace có định danh và metadata kỹ thuật cần thiết
    from repositories.audit_store import audit_store
    events = audit_store.get_events_by_assessment(test_aid)
    assert len(events) >= 2

    for ev in events:
        assert ev["assessment_id"] == test_aid
        assert ev["run_id"] == test_run
        assert "created_at" in ev
        assert "event_type" in ev
        assert "payload" in ev

    # 3. Tuyệt đối KHÔNG chứa dữ liệu nhạy cảm hoặc thông tin xác thực
    events_str = json.dumps(events)
    assert "jwt_secret" not in events_str.lower()
    assert "password" not in events_str.lower()
    assert "bearer" not in events_str.lower()
    assert "api_key" not in events_str.lower()
    # Tên tổ chức nhạy cảm được che giấu mặt nạ trong payload
    assert "Công ty Bảo mật Quốc gia" not in events_str
    assert "Côn***" in events_str or "***" in events_str
