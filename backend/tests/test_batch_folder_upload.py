"""Test Batch Folder Upload, Path Sanitization, XLSX Parser, Content Preview and Global Noise Pruning."""

import io
import os
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
import openpyxl


@pytest.fixture
def iso_client(tmp_path, monkeypatch):
    """FastAPI test client with isolated EVIDENCE_DIR."""
    from api.routes.iso27001 import router, EVIDENCE_DIR
    import api.routes.iso27001 as iso_mod

    test_ev_dir = str(tmp_path / "evidence")
    os.makedirs(test_ev_dir, exist_ok=True)
    monkeypatch.setattr(iso_mod, "EVIDENCE_DIR", test_ev_dir)

    app = FastAPI()
    app.include_router(router, prefix="/api")
    return TestClient(app), test_ev_dir


def test_batch_ingest_subfolder_path_no_errno2(iso_client):
    """Subfolder relative path from webkitdirectory must NOT trigger [Errno 2]."""
    client, ev_dir = iso_client

    # Create dummy docx content
    from docx import Document
    doc = Document()
    doc.add_heading("Bao Cao Danh Gia An Toan Thong Tin", level=1)
    doc.add_paragraph("He thong chua duoc cau hinh tuong lua day du. IP may chu: 10.140.0.3")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    # Upload with folder relative path
    subfolder_filename = "002.REPORT/PRJ_ENTERPRISE_BAO_CAO_VA_DOT4_2026.docx"
    response = client.post(
        "/api/iso27001/evidence/batch-ingest",
        files=[("files", (subfolder_filename, docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))]
    )
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "success"
    assert data["summary"]["processed_successfully"] == 1
    assert data["summary"]["failed_count"] == 0

    # Ensure no Errno 2 in errors
    assert len(data.get("errors", [])) == 0

    # Check file was saved cleanly
    proc_file = data["files"][0]
    assert proc_file["clean_name"] == "PRJ_ENTERPRISE_BAO_CAO_VA_DOT4_2026.docx"
    assert proc_file["status"] == "success"
    assert proc_file["char_count"] > 0

    # Verify report filename is NOT used as detected hostname
    detected = data.get("detected_hosts", [])
    for host in detected:
        assert "002.REPORT" not in (host.get("hostname") or "")
        assert "BAO CAO" not in (host.get("hostname") or "")


def test_xlsx_parser_extraction():
    """Verify openpyxl extracts tabular data across sheets with row/column structure."""
    from services.evidence_parser import parse_evidence_file

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Asset_Inventory"
    ws.append(["ID", "Asset Name", "Owner", "Criticality", "IP Address"])
    ws.append(["AST-01", "Core DB Server", "DBA Team", "High", "192.168.10.50"])
    ws.append(["AST-02", "Firewall Gateway", "NetOps", "Critical", "192.168.10.1"])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    res = parse_evidence_file(xlsx_bytes, "Danh_Muc_Tai_San_2026.xlsx")
    assert res["status"] == "success"
    assert res["char_count"] > 0
    assert "Core DB Server" in res["parsed_text"]
    assert "Firewall Gateway" in res["parsed_text"]
    assert "192.168.10.50" in res["parsed_text"]
    assert res["sha256"] != ""


def test_evidence_content_preview_and_delete(iso_client):
    """Test GET /evidence/file-content and DELETE /evidence/file/{filename}."""
    client, ev_dir = iso_client

    sample_content = b"Host Name: srv-app01\nOS Name: Ubuntu 22.04 LTS\nFirewall: Active iptables\nIP: 192.168.1.55"
    filename = "syslog_srv_app01.log"

    # Ingest
    res_ingest = client.post(
        "/api/iso27001/evidence/batch-ingest",
        files=[("files", (filename, sample_content, "text/plain"))]
    )
    assert res_ingest.status_code == 200

    # Preview content
    res_content = client.get(f"/api/iso27001/evidence/file-content?filename={filename}")
    assert res_content.status_code == 200
    content_data = res_content.json()
    assert content_data["status"] == "success"
    assert "srv-app01" in content_data["full_text"]
    assert content_data["sha256"] != ""

    # Delete globally
    res_delete = client.delete(f"/api/iso27001/evidence/file/{filename}")
    assert res_delete.status_code == 200
    del_data = res_delete.json()
    assert del_data["status"] == "success"
    assert len(del_data["removed_from"]) > 0

    # Verify content endpoint now 404s
    res_after_del = client.get(f"/api/iso27001/evidence/file-content?filename={filename}")
    assert res_after_del.status_code == 404
