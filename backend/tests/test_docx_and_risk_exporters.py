"""Tests for DOCX and Risk Register Exporters."""

import pytest
from services.report_docx_generator import generate_report_docx
from services.risk_register_exporter import generate_risk_register_xlsx


def test_generate_report_docx():
    sample_data = {
        "assessment_id": "test-assessment-docx-001",
        "run_id": "run-docx-001",
        "standard": "iso27001",
        "created_at": "2026-09-05T08:00:00Z",
        "compliance_percent": 74.5,
        "system_info": {
            "organization": {"name": "Tập đoàn Điện lực Việt Nam (EVN)", "industry": "Năng lượng"},
            "assessment_standard": "iso27001"
        },
        "json_data": {
            "assessment_id": "test-assessment-docx-001",
            "run_id": "run-docx-001",
            "compliance": {"percentage": 74.5},
            "weight_breakdown": {
                "critical": {"total": 20, "implemented": 15, "percent": 75.0},
                "high": {"total": 30, "implemented": 22, "percent": 73.3},
                "medium": {"total": 30, "implemented": 20, "percent": 66.7},
                "low": {"total": 13, "implemented": 12, "percent": 92.3},
            },
            "risk_register": [
                {
                    "control_id": "A.8.8",
                    "label": "Quản lý lỗ hổng kỹ thuật",
                    "gap": "Phát hiện hệ điều hành Windows Server 2008 R2 EOL và Hotfix N/A",
                    "severity": "critical",
                    "likelihood": 4,
                    "impact": 4,
                    "risk_score": 16,
                    "recommendation": "Nâng cấp OS lên Windows Server 2022 và cô lập phân vùng."
                }
            ]
        }
    }

    docx_bytes = generate_report_docx(sample_data)
    assert docx_bytes is not None
    assert len(docx_bytes) > 2000
    # Check PK zip header for docx
    assert docx_bytes[:2] == b"PK"


def test_generate_tcvn11930_capdo3_report_docx():
    tcvn_sample = {
        "assessment_id": "test-assessment-tcvn-001",
        "run_id": "run-tcvn-001",
        "standard": "tcvn11930",
        "created_at": "2026-09-08T08:00:00Z",
        "compliance_percent": 52.9,
        "system_info": {
            "organization": {"name": "Công ty TNHH MTV Nhiệt điện Thủ Đức (EVN TPC)", "industry": "Năng lượng"},
            "assessment_standard": "tcvn11930",
            "servers": 9,
            "firewalls": 2,
            "implemented_controls": [
                "NW.01", "NW.02", "NW.04", "NW.05",
                "SV.01", "SV.02", "SV.05",
                "APP.01", "APP.02", "APP.04", "APP.07",
                "DAT.01", "DAT.02", "DAT.03",
                "MNG.01", "MNG.02", "MNG.03", "MNG.04"
            ]
        },
        "json_data": {
            "assessment_id": "test-assessment-tcvn-001",
            "run_id": "run-tcvn-001",
            "compliance": {"percentage": 52.9},
            "risk_register": [
                {
                    "control_id": "SV.07",
                    "gap": "Hệ điều hành Windows Server 2008 R2 EOL và thiếu Hotfix KB5070247",
                    "severity": "critical",
                    "likelihood": 4,
                    "impact": 4,
                    "risk_score": 16,
                    "recommendation": "Lập kế hoạch nâng cấp HĐH và cập nhật bản vá khẩn cấp."
                },
                {
                    "control_id": "SV.03",
                    "gap": "Chưa cài đặt giải pháp EDR trên 9 máy chủ điều hành",
                    "severity": "high",
                    "likelihood": 4,
                    "impact": 4,
                    "risk_score": 16,
                    "recommendation": "Trang bị phần mềm EDR tập trung."
                }
            ]
        }
    }

    docx_bytes = generate_report_docx(tcvn_sample)
    assert docx_bytes is not None
    assert len(docx_bytes) > 2000
    assert docx_bytes[:2] == b"PK"


def test_generate_risk_register_xlsx():
    sample_data = {
        "assessment_id": "test-assessment-risk-001",
        "run_id": "run-risk-001",
        "standard": "ISO 27001:2022",
        "system_info": {"organization": {"name": "EVN Enterprise Infrastructure"}},
        "json_data": {
            "assessment_id": "test-assessment-risk-001",
            "run_id": "run-risk-001",
            "risk_register": [
                {
                    "control_id": "A.8.20",
                    "label": "Bảo mật mạng",
                    "category": "A.8 Công nghệ",
                    "gap": "Chưa kích hoạt IPS trên đường truyền kết nối trạm",
                    "severity": "high",
                    "likelihood": 4,
                    "impact": 4,
                    "risk_score": 16,
                    "recommendation": "Kích hoạt IPS profile và cập nhật signature hàng ngày."
                }
            ]
        }
    }

    xlsx_bytes = generate_risk_register_xlsx(assessment_data=sample_data)
    assert xlsx_bytes is not None
    assert len(xlsx_bytes) > 1000
    # Check PK zip header for xlsx
    assert xlsx_bytes[:2] == b"PK"


def test_api_export_endpoints(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    from main import app
    import api.routes.iso27001 as iso_mod

    # Setup dummy assessment
    test_id = "test_audit_001"
    sample_data = {
        "id": test_id,
        "standard": "iso27001",
        "created_at": "2026-09-05T08:00:00Z",
        "compliance_percent": 85.0,
        "system_info": {
            "organization": {"name": "Test Enterprise"},
            "assessment_standard": "iso27001"
        },
        "json_data": {
            "compliance": {"percentage": 85.0},
            "risk_register": [
                {"control_id": "A.5.1", "severity": "low", "likelihood": 2, "impact": 2, "risk_score": 4}
            ]
        }
    }

    monkeypatch.setattr(iso_mod, "load_assessment", lambda aid: sample_data if aid == test_id else None)

    client = TestClient(app)

    # Test DOCX endpoint
    resp_docx = client.post(f"/api/iso27001/assessments/{test_id}/export-docx")
    assert resp_docx.status_code == 200
    assert resp_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert resp_docx.content[:2] == b"PK"

    # Test Risk Register endpoint
    resp_risk = client.post(f"/api/iso27001/assessments/{test_id}/export-risk-register")
    assert resp_risk.status_code == 200
    assert resp_risk.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp_risk.content[:2] == b"PK"

    # Test PDF endpoint with weasyprint & healed report
    sample_data_completed = dict(sample_data)
    sample_data_completed["status"] = "completed"
    sample_data_completed["result"] = {
        "report": (
            "## 1. ĐÁNH GIÁ TỔNG QUAN\nTuân thủ: 85%\n\n"
            "## 5. EXECUTIVE SUMMARY\n"
            "a) Metrics Overview\n"
            "- Tỷ lệ Tuân thủ: 85%\n\n"
            "### b) Top 3 R"
        ),
        "json_data": sample_data["json_data"],
        "compliance_percent": 85.0
    }
    monkeypatch.setattr(iso_mod, "load_assessment", lambda aid: sample_data_completed if aid == test_id else None)

    resp_pdf = client.post(f"/api/iso27001/assessments/{test_id}/export-pdf")
    assert resp_pdf.status_code == 200
    # Returns PDF or fallback HTML
    assert resp_pdf.headers["content-type"] in ["application/pdf", "text/html; charset=utf-8"]
    if resp_pdf.headers["content-type"] == "application/pdf":
        assert resp_pdf.content[:4] == b"%PDF"


def test_chat_service_ensure_complete_report_heals_truncated():
    from services.chat_service import ChatService

    # Case 1: Truncated at '### b) Top 3 R'
    truncated_input = (
        "## 1. ĐÁNH GIÁ TỔNG QUAN\n"
        "Nội dung tổng quan...\n\n"
        "## 5. EXECUTIVE SUMMARY (TÓM TẮT CHO BAN LÃNH ĐẠO)\n"
        "a) Metrics Overview\n"
        "- Tỷ lệ Tuân thủ: 55.6%\n\n"
        "### b) Top 3 R"
    )

    healed = ChatService.ensure_complete_report(
        markdown_report=truncated_input,
        percentage=55.6,
        score=47,
        max_score=93,
        org_name="Công ty TNHH MTV Nhiệt điện Thủ Đức",
        std_name="ISO 27001:2022"
    )

    assert "### b) Top 3 R" not in healed.split("\n\n")[-1]
    assert "Top 3 Rủi ro trọng yếu" in healed
    assert "VND" in healed
    assert "Next Steps" in healed or "Lộ trình triển khai" in healed
    assert "30 ngày" in healed

    # Case 2: Complete report should not be damaged
    complete_input = (
        "## 1. ĐÁNH GIÁ TỔNG QUAN\n"
        "Nội dung...\n\n"
        "## 5. EXECUTIVE SUMMARY\n"
        "### b) Top 3 rủi ro:\n"
        "1. Rủi ro 1 rất dài và chi tiết có đầy đủ thông tin bảo mật và phân tích kỹ thuật hơn 150 ký tự để không bị coi là bị cắt ngắn giữa chừng khi kiểm tra.\n"
        "### c) Next Steps:\n"
        "1. Kế hoạch 30 ngày chi tiết."
    )
    res_complete = ChatService.ensure_complete_report(markdown_report=complete_input)
    assert res_complete == complete_input


def test_exporters_authoritative_weighted_compliance_regression():
    """Regression test: SoA, Risk Register, and DOCX must all use the exact backend WeightedCompliance."""
    import io
    import openpyxl
    from schemas.assessment_schema import (
        ControlItem,
        UnifiedAssessmentResult,
        WeightedCompliance,
    )
    from services.soa_exporter import generate_soa_xlsx
    from services.risk_register_exporter import generate_risk_register_xlsx
    from services.report_docx_generator import generate_report_docx

    aid = "asm_exporter_regression_001"
    run = "run_exporter_reg_001"
    controls = [
        ControlItem(
            control_id="A.5.1",
            label="Chính sách an toàn thông tin",
            weight="critical",
            assessment_verdict="satisfied",
            user_declaration="implemented",
            score=0,
            likelihood=2,
            impact=2,
            risk_score=4,
        ),
        ControlItem(
            control_id="A.5.2",
            label="Phân công trách nhiệm an toàn",
            weight="high",
            assessment_verdict="missing",
            user_declaration="not_implemented",
            score=0,
            gap="Chưa phân công trách nhiệm bảo mật",
            severity="high",
            likelihood=3,
            impact=3,
            risk_score=9,
        ),
    ]

    # Weighted Compliance: Critical(10)*1.0 + High(5)*0.0 = 10 / 15 = 66.7%
    backend_weighted = WeightedCompliance(
        weighted_score=10.0,
        weighted_max_score=15.0,
        percentage=66.7,
    )

    unified = UnifiedAssessmentResult(
        assessment_id=aid,
        run_id=run,
        standard="iso27001",
        controls=controls,
        weighted_compliance=backend_weighted,
        organization={"name": "Công ty Thẩm Định Độc Lập"},
    )

    # 1. SoA XLSX
    soa_bytes = generate_soa_xlsx(assessment_id=aid, assessment_data=unified)
    assert soa_bytes[:2] == b"PK"
    wb_soa = openpyxl.load_workbook(io.BytesIO(soa_bytes))
    ws_soa = wb_soa.active
    # Banner contains backend weighted compliance
    assert "66.7%" in str(ws_soa["A2"].value)
    assert aid in str(ws_soa["A2"].value)
    assert run in str(ws_soa["A2"].value)

    # 2. Risk Register XLSX
    rr_bytes = generate_risk_register_xlsx(assessment_id=aid, assessment_data=unified)
    assert rr_bytes[:2] == b"PK"
    wb_rr = openpyxl.load_workbook(io.BytesIO(rr_bytes))
    ws_rr = wb_rr.active
    assert "66.7%" in str(ws_rr["A2"].value)
    assert aid in str(ws_rr["A2"].value)
    assert run in str(ws_rr["A2"].value)

    # 3. DOCX Report
    docx_bytes = generate_report_docx(unified.model_dump())
    assert docx_bytes[:2] == b"PK"
    from docx import Document
    doc = Document(io.BytesIO(docx_bytes))
    doc_text = " ".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "66.7%" in doc_text
    assert aid in doc_text
    assert run in doc_text


def test_export_endpoints_vietnamese_org_name_and_content_disposition(monkeypatch):
    """Regression test: Org names with complex Vietnamese diacritics must not raise UnicodeEncodeError."""
    from fastapi.testclient import TestClient
    from main import app
    import api.routes.iso27001 as iso_mod
    import io
    import openpyxl
    from docx import Document

    test_id = "test_vn_org_001"
    vn_data = {
        "id": test_id,
        "assessment_id": test_id,
        "status": "completed",
        "standard": "iso27001",
        "created_at": "2026-09-20T10:00:00Z",
        "system_info": {
            "organization": {"name": "Công ty TNHH MTV Hạ tầng Năng lượng An Phú"}
        },
        "json_data": {
            "assessment_id": test_id,
            "organization": {"name": "Công ty TNHH MTV Hạ tầng Năng lượng An Phú"},
            "compliance": {"percentage": 80.0},
            "risk_register": [
                {
                    "control_id": "A.5.1",
                    "label": "Chính sách ATTT",
                    "severity": "medium",
                    "likelihood": 2,
                    "impact": 2,
                    "risk_score": 4,
                    "gap": "Chưa phê duyệt định kỳ hàng năm",
                    "recommendation": "Phê duyệt lại quy chế ATTT."
                }
            ]
        }
    }
    monkeypatch.setattr(iso_mod, "load_assessment", lambda aid: vn_data if aid == test_id else None)
    client = TestClient(app)

    # 1. DOCX
    resp_docx = client.post(f"/api/iso27001/assessments/{test_id}/export-docx")
    assert resp_docx.status_code == 200
    assert resp_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    cd_docx = resp_docx.headers["content-disposition"]
    assert "filename=" in cd_docx
    assert "filename*=UTF-8''" in cd_docx
    doc = Document(io.BytesIO(resp_docx.content))
    assert len(doc.paragraphs) > 0

    # 2. Risk Register XLSX
    resp_risk = client.post(f"/api/iso27001/assessments/{test_id}/export-risk-register")
    assert resp_risk.status_code == 200
    assert resp_risk.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    cd_risk = resp_risk.headers["content-disposition"]
    assert "filename=" in cd_risk
    assert "filename*=UTF-8''" in cd_risk
    wb = openpyxl.load_workbook(io.BytesIO(resp_risk.content))
    assert wb.active is not None


def test_export_empty_risk_register_graceful():
    """Generator and route must handle empty risk/GAP lists without crashing and show required notice."""
    import io
    import openpyxl
    from docx import Document
    from schemas.assessment_schema import ControlItem, UnifiedAssessmentResult, WeightedCompliance

    controls_satisfied = [
        ControlItem(
            control_id="A.5.1",
            label="Chính sách an toàn thông tin",
            weight="critical",
            assessment_verdict="satisfied",
            user_declaration="implemented",
            score=0,
            likelihood=1,
            impact=1,
            risk_score=1,
        )
    ]
    unified = UnifiedAssessmentResult(
        assessment_id="test_all_satisfied_001",
        run_id="run_satisfied_001",
        standard="iso27001",
        status="completed",
        controls=controls_satisfied,
        risk_register=[],
        top_gaps=[],
        weighted_compliance=WeightedCompliance(weighted_score=10.0, weighted_max_score=10.0, percentage=100.0),
        organization={"name": "Công ty Không Có Rủi Ro"},
    )

    # 1. Risk Register XLSX
    rr_bytes = generate_risk_register_xlsx(assessment_id="test_all_satisfied_001", assessment_data=unified)
    assert rr_bytes[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(rr_bytes))
    ws = wb.active
    # Must display the required notice on row 4
    assert "Không có mục rủi ro được tạo từ kết quả hiện tại" in str(ws["A4"].value)

    # 2. DOCX Report
    docx_bytes = generate_report_docx(unified.model_dump())
    assert docx_bytes[:2] == b"PK"
    doc = Document(io.BytesIO(docx_bytes))
    doc_text = " ".join([p.text for p in doc.paragraphs])
    assert "Không có mục rủi ro được tạo từ kết quả hiện tại" in doc_text


def test_export_nonexistent_assessment_returns_404():
    """Requesting an unrecorded assessment ID must return structured 404."""
    from fastapi.testclient import TestClient
    from main import app

    client = TestClient(app)
    resp_docx = client.post("/api/iso27001/assessments/non-existent-aid-99999/export-docx")
    assert resp_docx.status_code == 404
    data_docx = resp_docx.json()
    assert "Assessment not found" in (data_docx.get("detail") or data_docx.get("error") or "")

    resp_risk = client.post("/api/iso27001/assessments/non-existent-aid-99999/export-risk-register")
    assert resp_risk.status_code == 404
    data_risk = resp_risk.json()
    assert "Assessment not found" in (data_risk.get("detail") or data_risk.get("error") or "")


def test_export_incomplete_assessment_returns_400(monkeypatch):
    """Assessments with status not 'completed' must return structured 400."""
    from fastapi.testclient import TestClient
    from main import app
    import api.routes.iso27001 as iso_mod

    test_id = "test_processing_001"
    processing_data = {
        "id": test_id,
        "assessment_id": test_id,
        "status": "processing",
        "standard": "iso27001",
        "json_data": {"compliance": {"percentage": 10.0}}
    }
    monkeypatch.setattr(iso_mod, "load_assessment", lambda aid: processing_data if aid == test_id else None)
    client = TestClient(app)

    resp_docx = client.post(f"/api/iso27001/assessments/{test_id}/export-docx")
    assert resp_docx.status_code == 400
    err_docx = resp_docx.json().get("detail") or resp_docx.json().get("error") or ""
    assert "not completed yet" in err_docx

    resp_risk = client.post(f"/api/iso27001/assessments/{test_id}/export-risk-register")
    assert resp_risk.status_code == 400
    err_risk = resp_risk.json().get("detail") or resp_risk.json().get("error") or ""
    assert "not completed yet" in err_risk


def test_live_assessment_163f3f6b_exports_all_artefacts():
    """Verify all artefacts for real assessment 163f3f6b using both full ID and short ID."""
    from fastapi.testclient import TestClient
    from main import app
    import io
    import openpyxl
    from docx import Document

    client = TestClient(app)

    test_cases = [
        ("163f3f6b-1265-41ac-acab-10f9d7580da5", "163f3f6b", False),
        ("e9bee7ff-8f23-4f01-9666-7e271e7b1a77", "e9bee7ff", True),
    ]

    for full_id, short_id, is_tcvn in test_cases:
        for aid in [full_id, short_id]:
            # 1. DOCX
            resp_docx = client.post(f"/api/iso27001/assessments/{aid}/export-docx")
            assert resp_docx.status_code == 200, f"DOCX export failed for {aid}: {resp_docx.text}"
            assert resp_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            assert f"IT_Audit_Report_{short_id}.docx" in resp_docx.headers["content-disposition"]
            doc = Document(io.BytesIO(resp_docx.content))
            assert len(doc.paragraphs) > 10

            # 2. Risk Register XLSX
            resp_risk = client.post(f"/api/iso27001/assessments/{aid}/export-risk-register")
            assert resp_risk.status_code == 200, f"Risk Register export failed for {aid}: {resp_risk.text}"
            assert resp_risk.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            assert f"Risk_Register_{short_id}.xlsx" in resp_risk.headers["content-disposition"]
            wb_risk = openpyxl.load_workbook(io.BytesIO(resp_risk.content))
            assert wb_risk.active is not None

            # 3. SoA XLSX
            resp_soa = client.post("/api/iso27001/soa/export", json={"assessment_id": aid, "org_name": "Test Org"})
            assert resp_soa.status_code == 200, f"SoA export failed for {aid}: {resp_soa.text}"
            expected_soa_fn = f"SoA_TCVN11930_{short_id}.xlsx" if is_tcvn else f"SoA_ISO27001_{short_id}.xlsx"
            assert expected_soa_fn in resp_soa.headers["content-disposition"]

            # 4. PDF
            resp_pdf = client.post(f"/api/iso27001/assessments/{aid}/export-pdf")
            assert resp_pdf.status_code == 200
            assert resp_pdf.headers["content-type"] in ["application/pdf", "text/html; charset=utf-8"]
            assert f"Audit_Report_{short_id}.pdf" in resp_pdf.headers["content-disposition"]

            # 5. Assessment JSON (GET /api/iso27001/assessments/{id})
            resp_json = client.get(f"/api/iso27001/assessments/{aid}")
            assert resp_json.status_code == 200
            assert resp_json.json().get("id") or resp_json.json().get("assessment_id")

            # 6. Audit Trace JSON (GET /api/iso27001/assessments/{id}/audit-trace)
            resp_trace = client.get(f"/api/iso27001/assessments/{aid}/audit-trace")
            assert resp_trace.status_code == 200
            assert "events" in resp_trace.json() or "summary" in resp_trace.json()



