"""Mandatory Test Suite for Risk Heatmap (4×4 Matrix).

Validates:
1. All 16 combinations (Likelihood 1..4 × Impact 1..4)
2. Rejection of out-of-bounds inputs: likelihood=0, 5; impact=0, 5
3. 4×4 = 16 is maximum score; 20 and 25 cannot be generated
4. Pydantic schemas (ControlItem, RiskCreate, RiskUpdate) reject invalid inputs
5. Artifact validator flags inconsistent risk scores or scores > 16
6. Exporters (XLSX, DOCX) never produce risk scores > 16
7. API routes enforce 4×4 matrix and 1-4 boundaries
8. Legacy assessment preservation mode allows reading marked historical data
"""

import io
import pytest
from pydantic import ValidationError
import openpyxl
from docx import Document

from services.risk_scoring import (
    calculate_risk_score,
    classify_risk_severity,
    is_valid_risk_axis,
    is_valid_risk_score,
    RISK_AXIS_MIN,
    RISK_AXIS_MAX,
    RISK_SCORE_MIN,
    RISK_SCORE_MAX,
)
from schemas.assessment_schema import ControlItem, UnifiedAssessmentResult
from api.schemas.risk import RiskCreate, RiskUpdate, Risk
from services.artifact_validator import validate_assessment_artifacts
from services.risk_register_exporter import generate_risk_register_xlsx
from services.report_docx_generator import generate_report_docx


# ── 1. Test All 16 Combinations ───────────────────────────────────────────────

@pytest.mark.parametrize("likelihood", [1, 2, 3, 4])
@pytest.mark.parametrize("impact", [1, 2, 3, 4])
def test_all_16_risk_combinations(likelihood, impact):
    """Test every cell in the 4x4 matrix satisfies score = L * I and score <= 16."""
    score = calculate_risk_score(likelihood, impact)
    assert score == likelihood * impact
    assert RISK_SCORE_MIN <= score <= RISK_SCORE_MAX
    assert score <= 16

    # Verify severity classification is valid
    sev = classify_risk_severity(score)
    assert sev in ("critical", "high", "medium", "low")

    # Verify ControlItem schema accepts this valid combination
    item = ControlItem(
        control_id=f"CTRL_{likelihood}_{impact}",
        likelihood=likelihood,
        impact=impact,
        risk_score=score,
        risk_severity=sev,
    )
    assert item.likelihood == likelihood
    assert item.impact == impact
    assert item.risk_score == score


# ── 2. Boundary Rejections in Core Engine ─────────────────────────────────────

def test_risk_scoring_rejects_likelihood_0():
    with pytest.raises(ValueError, match="Likelihood must be between 1 and 4"):
        calculate_risk_score(0, 3)


def test_risk_scoring_rejects_likelihood_5():
    with pytest.raises(ValueError, match="Likelihood must be between 1 and 4"):
        calculate_risk_score(5, 3)


def test_risk_scoring_rejects_impact_0():
    with pytest.raises(ValueError, match="Impact must be between 1 and 4"):
        calculate_risk_score(3, 0)


def test_risk_scoring_rejects_impact_5():
    with pytest.raises(ValueError, match="Impact must be between 1 and 4"):
        calculate_risk_score(3, 5)


def test_maximum_score_is_16():
    """Verify maximum possible score is 16 and cannot produce 20 or 25."""
    max_score = calculate_risk_score(4, 4)
    assert max_score == 16
    assert max_score <= 16
    assert max_score != 20
    assert max_score != 25


# ── 3. Schema Boundary Rejections ────────────────────────────────────────────

def test_control_item_rejects_out_of_bounds():
    # likelihood = 5
    with pytest.raises(ValidationError):
        ControlItem(control_id="A.5.1", likelihood=5, impact=4, risk_score=20)

    # likelihood = 0
    with pytest.raises(ValidationError):
        ControlItem(control_id="A.5.1", likelihood=0, impact=4, risk_score=0)

    # impact = 5
    with pytest.raises(ValidationError):
        ControlItem(control_id="A.5.1", likelihood=4, impact=5, risk_score=20)

    # impact = 0
    with pytest.raises(ValidationError):
        ControlItem(control_id="A.5.1", likelihood=4, impact=0, risk_score=0)

    # Mismatched risk score
    with pytest.raises(ValidationError):
        ControlItem(control_id="A.5.1", likelihood=3, impact=3, risk_score=10)

    # Risk score > 16
    with pytest.raises(ValidationError):
        ControlItem(control_id="A.5.1", likelihood=4, impact=4, risk_score=25)


def test_risk_api_schemas_reject_out_of_bounds():
    payload_base = {
        "asset_ref": "Core Database",
        "threat": "SQL Injection",
        "vulnerability": "Lack of input sanitization",
        "treatment": "mitigate",
        "residual_score": 8,
        "owner": "SecOps",
        "review_date": "2026-12-31",
    }

    # Likelihood 5
    with pytest.raises(ValidationError):
        RiskCreate(**{**payload_base, "likelihood": 5, "impact": 4})

    # Likelihood 0
    with pytest.raises(ValidationError):
        RiskCreate(**{**payload_base, "likelihood": 0, "impact": 4})

    # Impact 5
    with pytest.raises(ValidationError):
        RiskCreate(**{**payload_base, "likelihood": 4, "impact": 5})

    # Impact 0
    with pytest.raises(ValidationError):
        RiskCreate(**{**payload_base, "likelihood": 4, "impact": 0})

    # Residual score 20 or 25
    with pytest.raises(ValidationError):
        RiskCreate(**{**payload_base, "likelihood": 4, "impact": 4, "residual_score": 20})

    with pytest.raises(ValidationError):
        RiskCreate(**{**payload_base, "likelihood": 4, "impact": 4, "residual_score": 25})

    # RiskUpdate also rejects
    with pytest.raises(ValidationError):
        RiskUpdate(likelihood=5)

    with pytest.raises(ValidationError):
        RiskUpdate(impact=5)

    with pytest.raises(ValidationError):
        RiskUpdate(residual_score=25)


# ── 4. Validator Detection of Inconsistencies ────────────────────────────────

def test_artifact_validator_detects_risk_inconsistencies():
    # 1. Valid data passes
    valid_assessment = {
        "assessment_id": "asm_valid_test",
        "run_id": "run_valid_01",
        "controls": [
            {
                "control_id": "A.5.1",
                "assessment_verdict": "missing",
                "likelihood": 4,
                "impact": 4,
                "risk_score": 16,
            }
        ],
        "risk_register": [
            {
                "control_id": "A.5.1",
                "likelihood": 4,
                "impact": 4,
                "risk_score": 16,
            }
        ],
    }
    res_valid = validate_assessment_artifacts("asm_valid_test", assessment_data=valid_assessment)
    assert "risk_score_consistency" in res_valid["checks"]
    assert res_valid["checks"]["risk_score_consistency"]["status"] == "PASS"

    # 2. Invalid likelihood = 5 detected
    invalid_l5 = {
        "assessment_id": "asm_invalid_l5",
        "run_id": "run_01",
        "controls": [
            {"control_id": "A.5.1", "likelihood": 5, "impact": 4, "risk_score": 20}
        ],
    }
    res_l5 = validate_assessment_artifacts("asm_invalid_l5", assessment_data=invalid_l5)
    assert res_l5["checks"]["risk_score_consistency"]["status"] == "FAIL"

    # 3. Invalid score != L * I detected
    invalid_math = {
        "assessment_id": "asm_invalid_math",
        "run_id": "run_01",
        "controls": [
            {"control_id": "A.5.1", "likelihood": 3, "impact": 3, "risk_score": 12}
        ],
    }
    res_math = validate_assessment_artifacts("asm_invalid_math", assessment_data=invalid_math)
    assert res_math["checks"]["risk_score_consistency"]["status"] == "FAIL"

    # 4. Invalid score > 16 detected
    invalid_max = {
        "assessment_id": "asm_invalid_max",
        "run_id": "run_01",
        "risk_register": [
            {"control_id": "A.5.1", "likelihood": 4, "impact": 4, "risk_score": 20}
        ],
    }
    res_max = validate_assessment_artifacts("asm_invalid_max", assessment_data=invalid_max)
    assert res_max["checks"]["risk_score_consistency"]["status"] == "FAIL"


# ── 5. Exporters Never Produce Scores > 16 ────────────────────────────────────

def test_risk_register_xlsx_exporter_bounds():
    sample_data = {
        "assessment_id": "asm_export_test",
        "run_id": "run_export_01",
        "standard": {"id": "iso27001", "name": "ISO/IEC 27001:2022"},
        "json_data": {
            "assessment_id": "asm_export_test",
            "run_id": "run_export_01",
            "risk_register": [
                {
                    "control_id": "A.5.1",
                    "severity": "critical",
                    "likelihood": 4,
                    "impact": 4,
                    "risk_score": 16,
                },
                {
                    "control_id": "A.8.8",
                    "severity": "high",
                    "likelihood": 3,
                    "impact": 3,
                    "risk_score": 9,
                },
            ]
        },
    }
    xlsx_bytes = generate_risk_register_xlsx(assessment_data=sample_data)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    # Check header indicates 1-4
    h_likelihood = ws.cell(row=3, column=7).value
    h_impact = ws.cell(row=3, column=8).value
    assert "1-4" in str(h_likelihood)
    assert "1-4" in str(h_impact)
    assert "1-5" not in str(h_likelihood)
    assert "1-5" not in str(h_impact)

    # Check data rows
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is not None:
            l_val = row[6]  # col 7 (0-indexed 6)
            i_val = row[7]  # col 8 (0-indexed 7)
            score_val = row[8]  # col 9 (0-indexed 8)
            assert 1 <= l_val <= 4
            assert 1 <= i_val <= 4
            assert score_val <= 16
            assert score_val == l_val * i_val


def test_docx_report_exporter_bounds():
    sample_data = {
        "assessment_id": "asm_docx_test",
        "run_id": "run_docx_01",
        "standard": {"id": "iso27001", "name": "ISO/IEC 27001:2022"},
        "json_data": {
            "assessment_id": "asm_docx_test",
            "run_id": "run_docx_01",
            "top_gaps": [
                {
                    "id": "A.5.1",
                    "severity": "critical",
                    "gap": "Test critical gap",
                }
            ]
        },
    }
    docx_bytes = generate_report_docx(sample_data)
    doc = Document(io.BytesIO(docx_bytes))
    full_text = "\n".join([p.text for p in doc.paragraphs])
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text

    assert "Likelihood: 1-4" in full_text
    assert "Impact: 1-4" in full_text
    assert "Likelihood: 1-5" not in full_text
    assert "Impact: 1-5" not in full_text


# ── 6. Legacy Data Compatibility ─────────────────────────────────────────────

def test_legacy_data_reading_compatibility():
    """Historical data with 5x5 metrics can be read when marked as legacy."""
    legacy_item = ControlItem(
        control_id="A.5.1",
        likelihood=5,
        impact=5,
        risk_score=25,
        is_legacy=True,
    )
    assert legacy_item.likelihood == 5
    assert legacy_item.impact == 5
    assert legacy_item.risk_score == 25
    assert legacy_item.is_legacy is True

    # Legacy Risk record
    legacy_risk = Risk(
        id="legacy_risk_123",
        asset_ref="Legacy Server",
        threat="Old Exploit",
        vulnerability="EOL Software",
        likelihood=5,
        impact=5,
        treatment="mitigate",
        residual_score=25,
        owner="Admin",
        review_date="2024-01-01",
        inherent_score=25,
        created_at="2024-01-01T00:00:00Z",
        updated_at="2024-01-01T00:00:00Z",
        is_legacy=True,
    )
    assert legacy_risk.likelihood == 5
    assert legacy_risk.inherent_score == 25
    assert legacy_risk.is_legacy is True
